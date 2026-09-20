# -*- coding: utf-8 -*-
"""
Tests del adelgazado del PPTX.

El deck pesaba 40 MB, de los que el 99,2% eran 82 imágenes heredadas de la
plantilla corporativa (un curso de 63 láminas sobre Joule). La propuesta usa 11.

Lo que se comprueba aquí no es solo que pese menos: es que **cada lámina
conserva exactamente el mismo contenido**. Un deck ligero pero con el fondo de
ondas o el logo rotos sería mucho peor que uno grande.
"""
import hashlib
import os
import re
import shutil
import zipfile

import pytest

from services.financial_engine import calculate_financials
from services.ppt_generator import generate_deck
from services.pptx_slim import adelgazar, validar_paquete, _relaciones

# Tope operativo. La mayoría de servidores de correo corporativos rechazan
# adjuntos por encima de 10-25 MB: por encima de esto, la propuesta deja de
# poder enviarse al cliente, que es justo para lo que existe.
MAX_MB_ENVIABLE = 10


@pytest.fixture(scope='module')
def financieros():
    return calculate_financials(['FI', 'CO', 'MM', 'SD', 'PP', 'PS'])


@pytest.fixture(scope='module')
def par_de_decks(tmp_path_factory, financieros):
    """
    El mismo deck generado dos veces: uno sin podar y otro podado.

    El generador ya adelgaza, así que para tener un original de referencia hay
    que desactivar ese paso: si no, se compararían dos decks ya podados y el
    test no demostraría nada.
    """
    from unittest.mock import patch
    carpeta = tmp_path_factory.mktemp('equivalencia')
    original = str(carpeta / 'original.pptx')
    podado = str(carpeta / 'podado.pptx')

    with patch('services.ppt_generator.adelgazar', lambda ruta: None):
        generate_deck("Equivalencia S.A.", "Minería y Recursos", "d", "Alta",
                      financieros, original, edition="Public")

    shutil.copy2(original, podado)
    adelgazar(podado)
    return original, podado


@pytest.fixture(scope='module')
def deck(tmp_path_factory, financieros):
    ruta = str(tmp_path_factory.mktemp('peso') / 'deck.pptx')
    generate_deck("Peso Test S.A.", "Minería y Recursos", "desc", "Alta",
                  financieros, ruta, edition="Private", theme='cobre')
    return ruta


class TestPesoDelEntregable:
    def test_el_deck_se_puede_enviar_por_correo(self, deck):
        mb = os.path.getsize(deck) / 1048576
        assert mb < MAX_MB_ENVIABLE, (
            f"El deck pesa {mb:.1f} MB: por encima de {MAX_MB_ENVIABLE} MB no se "
            f"puede adjuntar en un correo corporativo"
        )

    def test_solo_quedan_las_imagenes_necesarias(self, deck):
        with zipfile.ZipFile(deck) as z:
            imagenes = [n for n in z.namelist() if n.startswith('ppt/media/')]
        assert len(imagenes) < 25, (
            f"{len(imagenes)} imágenes: la plantilla trae 82 y la propuesta usa ~11"
        )

    def test_solo_quedan_los_layouts_usados(self, deck):
        with zipfile.ZipFile(deck) as z:
            layouts = [n for n in z.namelist()
                       if re.match(r'ppt/slideLayouts/slideLayout\d+\.xml$', n)]
        assert len(layouts) <= 10, f"{len(layouts)} layouts; la plantilla trae 122"


class TestElDeckSigueIntacto:
    """
    Podar partes obliga a rehacer presentation.xml, los sldLayoutIdLst de cada
    master y [Content_Types].xml. Si algo se queda a medias, PowerPoint da el
    archivo por corrupto.
    """

    def test_el_paquete_es_coherente(self, deck):
        ok, problemas = validar_paquete(deck)
        assert ok, problemas[:5]

    def test_el_zip_no_esta_corrupto(self, deck):
        with zipfile.ZipFile(deck) as z:
            assert z.testzip() is None

    def test_python_pptx_lo_reabre_con_todas_las_laminas(self, deck):
        from pptx import Presentation
        assert len(Presentation(deck).slides) == 12

    def test_los_graficos_nativos_sobreviven(self, deck):
        from pptx import Presentation
        graficos = sum(1 for s in Presentation(deck).slides
                       for sh in s.shapes if sh.has_chart)
        assert graficos == 2, "el gráfico de la lámina económica y el del ROI"

    def test_cada_lamina_conserva_su_layout(self, deck):
        with zipfile.ZipFile(deck) as z:
            nombres = set(z.namelist())
            laminas = [n for n in nombres if re.match(r'ppt/slides/slide\d+\.xml$', n)]
            for lamina in laminas:
                destinos = [d for _, d in _relaciones(z, lamina, nombres)]
                assert any('/slideLayouts/' in d for d in destinos), lamina

    def test_cada_layout_conserva_su_master(self, deck):
        with zipfile.ZipFile(deck) as z:
            nombres = set(z.namelist())
            layouts = [n for n in nombres
                       if re.match(r'ppt/slideLayouts/slideLayout\d+\.xml$', n)]
            for layout in layouts:
                destinos = [d for _, d in _relaciones(z, layout, nombres)]
                assert any('/slideMasters/' in d for d in destinos), layout

    def test_presentation_xml_no_referencia_masters_eliminados(self, deck):
        """Un sldMasterId huérfano hace que PowerPoint rechace el archivo."""
        with zipfile.ZipFile(deck) as z:
            nombres = set(z.namelist())
            ids_declarados = set(re.findall(
                r'<p:sldMasterId[^>]*r:id="([^"]+)"',
                z.read('ppt/presentation.xml').decode('utf-8', 'replace')))
            ids_con_destino = {i for i, d in _relaciones(z, 'ppt/presentation.xml', nombres)
                               if '/slideMasters/' in d}
        assert ids_declarados <= ids_con_destino, (
            f"sldMasterId sin relación: {ids_declarados - ids_con_destino}")

    def test_ningun_master_referencia_layouts_eliminados(self, deck):
        with zipfile.ZipFile(deck) as z:
            nombres = set(z.namelist())
            for master in (n for n in nombres
                           if re.match(r'ppt/slideMasters/slideMaster\d+\.xml$', n)):
                declarados = set(re.findall(
                    r'<p:sldLayoutId[^>]*r:id="([^"]+)"',
                    z.read(master).decode('utf-8', 'replace')))
                con_destino = {i for i, d in _relaciones(z, master, nombres)
                               if '/slideLayouts/' in d}
                assert declarados <= con_destino, (
                    f"{master}: sldLayoutId sin relación {declarados - con_destino}")


class TestEquivalenciaVisual:
    """
    La comprobación que de verdad importa: adelgazar no puede cambiar lo que
    el cliente ve. Se genera el mismo deck dos veces —uno intacto y otro
    podado— y se compara el XML de cada lámina y las imágenes que su cadena
    lámina → layout → master referencia.
    """

    @staticmethod
    def _huella(ruta):
        """Por lámina: hash de su XML y hashes de las imágenes que alcanza."""
        salida = {}
        with zipfile.ZipFile(ruta) as z:
            nombres = set(z.namelist())
            for lamina in sorted(n for n in nombres
                                 if re.match(r'ppt/slides/slide\d+\.xml$', n)):
                cadena = [lamina]
                layouts = [d for _, d in _relaciones(z, lamina, nombres)
                           if '/slideLayouts/' in d]
                cadena += layouts
                for layout in layouts:
                    cadena += [d for _, d in _relaciones(z, layout, nombres)
                               if '/slideMasters/' in d]
                imagenes = set()
                for parte in cadena:
                    for _, destino in _relaciones(z, parte, nombres):
                        if destino.startswith('ppt/media/'):
                            imagenes.add(hashlib.sha256(z.read(destino)).hexdigest())
                salida[lamina] = (hashlib.sha256(z.read(lamina)).hexdigest(),
                                  tuple(sorted(imagenes)))
        return salida

    def test_el_contenido_de_cada_lamina_es_identico(self, par_de_decks):
        original, podado = par_de_decks
        antes, despues = self._huella(original), self._huella(podado)
        assert set(antes) == set(despues), "cambió el conjunto de láminas"
        distintas = [k for k in antes if antes[k] != despues[k]]
        assert not distintas, f"láminas alteradas por el adelgazado: {distintas}"

    def test_la_reduccion_es_sustancial(self, par_de_decks):
        original, podado = par_de_decks
        antes = os.path.getsize(original)
        despues = os.path.getsize(podado)
        assert despues < antes * 0.25, (
            f"solo se redujo de {antes/1048576:.1f} a {despues/1048576:.1f} MB"
        )


class TestRobustez:
    def test_adelgazar_dos_veces_es_inocuo(self, tmp_path, financieros):
        """La segunda pasada no debe encontrar nada que quitar ni romper nada."""
        ruta = str(tmp_path / 'doble.pptx')
        generate_deck("Doble S.A.", "Servicios Comerciales", "d", "Media",
                      financieros, ruta)
        tam1 = os.path.getsize(ruta)
        resultado = adelgazar(ruta)
        assert resultado['partes_eliminadas'] == 0
        assert os.path.getsize(ruta) == tam1
        ok, problemas = validar_paquete(ruta)
        assert ok, problemas[:5]

    def test_un_archivo_ilegible_no_se_destruye(self, tmp_path):
        """Ante un error se conserva el original: mejor grande que corrupto."""
        falso = tmp_path / 'noesunpptx.pptx'
        falso.write_bytes(b'esto no es un paquete OOXML')
        resultado = adelgazar(str(falso))
        assert resultado['partes_eliminadas'] == 0
        assert falso.read_bytes() == b'esto no es un paquete OOXML'

    def test_un_archivo_inexistente_lanza(self, tmp_path):
        with pytest.raises(FileNotFoundError):
            adelgazar(str(tmp_path / 'no-existe.pptx'))
